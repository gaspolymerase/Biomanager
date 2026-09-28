"""Experiments on every database's animals (app/experiments.py): fish,
flies, worms and organisms as well as mice, each with a readout that fits
it, manipulations of its own kinds, and the same page."""
from __future__ import annotations

from datetime import date, timedelta

# tests.base first: it points the app at a throwaway database before app is imported.
from tests.base import AppTestCase, count, location, one, uniq
from app import experiments as xp  # noqa: E402

TODAY = date.today()


class Base(AppTestCase):
    def new(self, db_key, readout="", start=None, client=None):
        client = client or self.a
        r = client.post(f"/experiments/in/{db_key}/create", data={
            "name": uniq("Exp "), "readout": readout, "start_date": (start or TODAY - timedelta(days=2)).isoformat()})
        self.assertEqual(r.status_code, 302, r.get_data(as_text=True)[:300])
        return int(location(r).rsplit("/", 1)[1])

    def data(self, exp, client=None):
        return (client or self.a).get(f"/experiments/{exp}/data.json").get_json()

    def add(self, exp, how, value, group="", client=None):
        return (client or self.a).post(f"/experiments/{exp}/subjects/add", json={"how": how, "value": value, "group": group})

    def read(self, exp, on, values, client=None):
        return (client or self.a).post(f"/experiments/{exp}/readings", json={"on": on.isoformat(), "values": values})


class Zebrafish(Base):
    def setUp(self):
        super().setUp()
        self.tank = self.make_tank(self.a)
        for n in (12, 8):
            self.a.post("/zebrafish/fish/create", data={"tank_id_fk": str(self.tank), "count": str(n)})

    def test_a_tank_of_fish_counted_as_they_survive(self):
        exp = self.new("zebrafish")
        data = self.data(exp)
        self.assertEqual(data["readout"]["key"], "survival")                 # fish: survival first
        self.assertIn("immersion", [k["key"] for k in data["kinds"]])      # and their own manipulations
        self.assertNotIn("injection", [k["key"] for k in data["kinds"]])
        r = self.add(exp, "group", str(self.tank), group="Tricaine")
        self.assertEqual(r.status_code, 200, r.get_json())
        subs = r.get_json()["subjects"]
        self.assertEqual(sorted(s["start"] for s in subs), [8, 12])
        big = next(s["key"] for s in subs if s["start"] == 12)
        data = self.read(exp, TODAY, {big: "9"}).get_json()
        row = next(r for r in data["table"]["rows"] if r["key"] == big)
        self.assertEqual((row["values"], row["pct"]), ([9.0], [75.0]))
        self.assertEqual(data["table"]["days"], [3])                          # started two days ago
        bad = self.read(exp, TODAY, {big: "13"})
        self.assertEqual(bad.status_code, 400)
        self.assertIn("started with 12", bad.get_json()["error"])

    def test_the_page_and_the_list(self):
        exp = self.new("zebrafish")
        html = self.get_ok(self.a, f"/experiments/{exp}")
        self.assertIn("Record manipulation", html)
        self.assertIn("Survival", html)
        listing = self.get_ok(self.a, "/experiments/in/zebrafish")
        self.assertIn(f"/experiments/{exp}", listing)
        self.assertNotIn(f"/experiments/{exp}\"", self.get_ok(self.a, "/colony?view=experiments"))
        self.assertIn("/experiments/in/zebrafish", self.get_ok(self.a, "/zebrafish?view=fish"))


class FlyStocks(Base):
    def test_vials_with_a_drug_in_the_food_and_their_survival(self):
        key = self.make_stock_module(self.a)
        vials = [self.make_vial(self.a, key) for _ in range(2)]
        exp = self.new(f"stocks:{key}", start=TODAY)
        data = self.data(exp)
        self.assertEqual(data["readout"]["key"], "survival")
        self.assertIn("food", [k["key"] for k in data["kinds"]])
        for v in vials:
            self.add(exp, "one", str(v), group="Paraquat" if v == vials[0] else "Control")
        subs = self.data(exp)["subjects"]
        first = next(s["key"] for s in subs if s["group"] == "Paraquat")
        self.assertIsNone(subs[0]["start"])                                  # flies aren't counted in a vial
        self.a.post(f"/experiments/{exp}/subjects/{first}/update", json={"start": "20"})
        # Drug in the food, recorded as it happens: it joins the plan on day 1.
        r = self.a.post(f"/colony/experiments/{exp}/record-now", json={
            "kind": "food", "agent": "Paraquat", "dose": "10 mM", "group": "Paraquat", "done_on": TODAY.isoformat()})
        self.assertEqual(r.status_code, 200, r.get_json())
        done = r.get_json()["schedule"][0]
        self.assertEqual((done["day"], done["title"], done["record"]["count"]), (1, "Paraquat 10 mM", 1))
        pct = self.read(exp, TODAY, {first: "15"}).get_json()["table"]["rows"]
        self.assertEqual(next(r for r in pct if r["key"] == first)["pct"], [75.0])

    def test_worms_have_rnai_and_brood_size(self):
        key = self.make_stock_module(self.a, kind="worm")
        data = self.data(self.new(f"stocks:{key}", readout="brood"))
        self.assertEqual((data["readout"]["label"], data["readout"]["kind"]), ("Brood size", "value"))
        self.assertIn("rnai", [k["key"] for k in data["kinds"]])


class Organisms(Base):
    def test_an_organism_weighed_and_dosed_by_weight(self):
        key = self.make_organism_module(self.a)
        animal = self.make_animal(self.a, key)
        exp = self.new(f"organisms:{key}", start=TODAY)
        self.assertEqual(self.data(exp)["readout"]["key"], "body_weight")
        self.add(exp, "one", str(animal))
        subject = f"organism:{animal}"
        self.read(exp, TODAY, {subject: "30"})
        step = self.a.post(f"/colony/experiments/{exp}/steps/save", json={
            "kind": "injection", "agent": "Drug", "dose": "10 mg/kg", "days": "1"}).get_json()["steps"][0]["id"]
        detail = self.a.get(f"/colony/experiments/{exp}/steps/{step}/day/1").get_json()
        self.assertEqual(detail["animals"][0]["amount"], "0.3 mg")


class Mice(Base):
    def test_another_readout_for_mice_is_kept_with_the_experiment(self):
        r = self.a.post("/colony/experiments/create", data={"name": uniq("Tumour "), "start_date": TODAY.isoformat()})
        exp = int(location(r).rsplit("/", 1)[1])
        mouse = self.make_mouse(self.a, self.admin)
        self.add(exp, "one", str(mouse))
        self.a.post(f"/colony/experiments/{exp}/add-mouse", data={"mouse_row_id": mouse})   # the old way too: no duplicate
        self.assertEqual(count("experiment_mice", "experiment_id_fk=?", exp), 1)
        self.a.post(f"/experiments/{exp}/update", data={"readout": "tumour_volume"})
        self.read(exp, TODAY, {f"mouse:{mouse}": "112.5"})
        self.assertEqual(one("select value from experiment_readings where experiment_id_fk=?", exp), 112.5)
        self.assertEqual(count("mouse_weights", "mouse_id_fk=?", mouse), 0)
        # Body weight again: the mouse's own weights.
        self.a.post(f"/experiments/{exp}/update", data={"readout": "body_weight"})
        self.read(exp, TODAY, {f"mouse:{mouse}": "24.2"})
        self.assertEqual(one("select grams from mouse_weights where mouse_id_fk=?", mouse), 24.2)
        # A mouse experiment's generic address is the colony's page.
        self.assertEqual(location(self.a.get(f"/experiments/{exp}")), f"/colony/experiments/{exp}")

    def test_toggling_one_mouse_in_a_recorded_day(self):
        r = self.a.post("/colony/experiments/create", data={"name": uniq("TAM "), "start_date": TODAY.isoformat()})
        exp = int(location(r).rsplit("/", 1)[1])
        mice = [self.make_mouse(self.a, self.admin) for _ in range(2)]
        for m in mice:
            self.add(exp, "one", str(m))
        data = self.a.post(f"/colony/experiments/{exp}/record-now", json={
            "kind": "injection", "agent": "Tamoxifen", "dose": "20 mg/kg", "done_on": TODAY.isoformat()}).get_json()
        step = data["schedule"][0]["step_id"]
        self.assertEqual(data["schedule"][0]["record"]["count"], 2)
        data = self.a.post(f"/colony/experiments/{exp}/steps/{step}/day/1/subject",
                           json={"subject": f"mouse:{mice[0]}", "given": False}).get_json()
        self.assertEqual([e["subject"] for e in data["schedule"][0]["record"]["subjects"]], [f"mouse:{mice[1]}"])


class Rules(Base):
    def test_only_the_owner_or_an_admin_changes_it(self):
        exp = self.new("zebrafish", client=self.a)
        self.assertEqual(self.add(exp, "group", "1", client=self.m).status_code, 403)
        self.assertEqual(self.read(exp, TODAY, {}, client=self.m).status_code, 403)
        self.assertEqual(self.m.post(f"/experiments/{exp}/update", data={"name": "x"}).status_code, 403)
        self.assertEqual(self.m.get(f"/experiments/{exp}/data.json").status_code, 200)       # but may look

    def test_no_readout_ahead_and_the_page_needs_its_database(self):
        exp = self.new("zebrafish")
        self.assertEqual(self.read(exp, TODAY + timedelta(days=1), {}).status_code, 400)
        self.assertEqual(self.a.get("/experiments/in/stocks:no-such").status_code, 404)

    def test_deleting_leaves_the_animals(self):
        tank = self.make_tank(self.a)
        self.a.post("/zebrafish/fish/create", data={"tank_id_fk": str(tank), "count": "3"})
        exp = self.new("zebrafish")
        self.add(exp, "group", str(tank))
        before = count("fish")
        r = self.post(self.a, f"/experiments/{exp}/delete")
        self.assertIn("Deleted experiment", r.get_data(as_text=True))
        self.assertEqual((count("experiments", "id=?", exp), count("experiment_subjects", "experiment_id_fk=?", exp)), (0, 0))
        self.assertEqual(count("fish"), before)

    def test_readouts_and_kinds_fit_the_animal(self):
        self.assertEqual(xp.READOUT_CHOICES["fly"][0], "survival")
        self.assertNotIn("body_weight", xp.READOUT_CHOICES["fly"] + xp.READOUT_CHOICES["worm"])
        self.assertEqual(xp.kind_info("worm", "weigh"), ("Readout day", "chart"))


class Notebook(Base):
    def test_the_block_reads_any_experiment_and_its_readout(self):
        key = self.make_stock_module(self.a)
        vial = self.make_vial(self.a, key)
        exp = self.new(f"stocks:{key}", start=TODAY)
        self.add(exp, "one", str(vial))
        self.a.post(f"/experiments/{exp}/subjects/unit:{vial}/update", json={"start": "20"})
        self.read(exp, TODAY, {f"unit:{vial}": "17"})
        got = self.a.get(f"/colony/experiments/{exp}/notebook.json").get_json()["experiment"]
        self.assertEqual((got["readout"]["label"], got["nouns"]), ("Survival", got["nouns"]))
        self.assertEqual(got["weights"]["rows"][0]["pct"], [85.0])
        self.assertEqual(got["url"], f"/experiments/{exp}")
        listed = self.a.get("/colony/experiments/notebook-list.json").get_json()["experiments"]
        self.assertIn(exp, [e["id"] for e in listed])
        r = self.a.post(f"/colony/experiments/{exp}/notebook")
        body = one("select body from notebook_pages where id=?", int(location(r).split("page=")[1]))
        self.assertIn("```experiment", body)


class Regimens(Base):
    def mouse_exp(self, start=TODAY):
        r = self.a.post("/colony/experiments/create", data={"name": uniq("Asthma "), "start_date": start.isoformat()})
        return int(location(r).rsplit("/", 1)[1])

    def test_a_saved_regimen_plans_the_next_cohort_but_records_nothing(self):
        first = self.mouse_exp()
        self.a.post(f"/colony/experiments/{first}/steps/save", json={"kind": "injection", "agent": "Tamoxifen", "dose": "20 mg/kg", "days": "1"})
        self.a.post(f"/colony/experiments/{first}/steps/save", json={"kind": "challenge", "agent": "HDM", "dose": "25 µg", "days": "2-5"})
        name = uniq("TAM then HDM ")
        self.a.post(f"/colony/experiments/{first}/regimens/save", json={"name": name})
        second = self.mouse_exp()
        regimen = next(r for r in self.data(second)["regimens"] if r["name"] == name)
        data = self.a.post(f"/colony/experiments/{second}/regimens/{regimen['id']}/apply", json={}).get_json()
        self.assertEqual([(s["agent"], s["days_label"]) for s in data["steps"]], [("Tamoxifen", "1"), ("HDM", "2–5")])
        self.assertTrue(all(r["record"] is None for r in data["schedule"]))          # planned, not done
        self.assertEqual(count("experiment_step_records", "step_id_fk in (select id from experiment_steps where experiment_id_fk=?)", second), 0)

    def test_the_owner_is_reminded_once_a_day_of_what_is_due(self):
        from app import notify
        from app.db import SessionLocal
        from app.models import UserAccount
        from tests.base import client_for, make_user
        who = make_user(uniq("owner"))                 # nobody who opened a page today, so not reminded yet
        c = client_for(who)
        r = c.post("/colony/experiments/create", data={"name": uniq("Asthma "), "start_date": (TODAY - timedelta(days=1)).isoformat()})
        exp = int(location(r).rsplit("/", 1)[1])
        c.post(f"/colony/experiments/{exp}/steps/save", json={"kind": "challenge", "agent": "HDM", "dose": "25 µg", "days": "1-2"})
        execute_ = __import__("tests.base", fromlist=["execute"]).execute
        execute_("delete from notifications where recipient_username=?", who)
        from app.app import app
        with app.test_request_context(), SessionLocal() as s:
            from flask import g
            user = s.scalar(__import__("sqlalchemy").select(UserAccount).where(UserAccount.username == who))
            g.user = user
            self.assertTrue(notify.daily_experiment_reminder(s, user))
            s.commit()
            self.assertFalse(notify.daily_experiment_reminder(s, user))               # once a day
        title = one("select title from notifications where recipient_username=? and category='experiments' order by id desc limit 1", who)
        self.assertIn("HDM 25 µg", title)
        self.assertIn("overdue", title)


class ReagentsAndSamples(Base):
    def test_the_lot_used_is_kept_and_an_expired_one_is_flagged(self):
        item = self.make_item(self.a, "reagents", name=uniq("HDM extract "), lot="L-77", expires_on=(TODAY - timedelta(days=3)).isoformat())
        r = self.a.post("/colony/experiments/create", data={"name": uniq("HDM "), "start_date": TODAY.isoformat()})
        exp = int(location(r).rsplit("/", 1)[1])
        self.add(exp, "one", str(self.make_mouse(self.a, self.admin)))
        self.assertIn(item, [x["id"] for x in self.data(exp)["reagents"]])
        data = self.a.post(f"/colony/experiments/{exp}/record-now", json={
            "kind": "challenge", "agent": "HDM", "dose": "25 µg", "done_on": TODAY.isoformat(), "reagent_id": str(item)}).get_json()
        rec = data["schedule"][0]["record"]
        self.assertEqual((rec["reagent"]["lot"], rec["reagent"]["id"]), ("L-77", item))
        self.assertTrue(any("expired" in p for p in data["problems"]))

    def test_sample_records_only_when_asked(self):
        r = self.a.post("/colony/experiments/create", data={"name": uniq("Harvest "), "start_date": TODAY.isoformat()})
        exp = int(location(r).rsplit("/", 1)[1])
        mice = [self.make_mouse(self.a, self.admin) for _ in range(2)]
        for m in mice:
            self.add(exp, "one", str(m))
        self.assertIn("samples", [i["key"] for i in self.data(exp)["sample_inventories"]])
        before = count("inventory_items")
        self.a.post(f"/colony/experiments/{exp}/record-now", json={"kind": "sample", "agent": "Blood", "done_on": TODAY.isoformat()})
        self.assertEqual(count("inventory_items"), before)                             # not asked: none
        data = self.a.post(f"/colony/experiments/{exp}/record-now", json={
            "kind": "sample", "agent": "Lung", "done_on": TODAY.isoformat(), "make_samples": True,
            "sample_inventory": "samples", "sample_name": "Lung"}).get_json()
        made = data["schedule"][-1]["record"]["samples"]
        self.assertEqual(len(made), 2)
        mouse_id = one("select mouse_id from mice where id=?", mice[0])
        attrs = __import__("json").loads(one("select attrs from inventory_items where id=?", made[0]["id"]))
        self.assertEqual(attrs["source"], {"kind": "mouse", "ref": str(mouse_id)})
        self.assertIn("Lung", one("select name from inventory_items where id=?", made[0]["id"]))


class StatsAndExport(Base):
    def test_groups_compared_day_by_day_and_exported(self):
        from app import exp_stats
        tank_a, tank_b = self.make_tank(self.a), self.make_tank(self.a)
        for t in (tank_a, tank_b):
            for _ in range(3):
                self.a.post("/zebrafish/fish/create", data={"tank_id_fk": str(t), "count": "10"})
        exp = self.new("zebrafish", readout="length", start=TODAY)
        self.add(exp, "group", str(tank_a), group="Drug")
        self.add(exp, "group", str(tank_b), group="Vehicle")
        subs = self.data(exp)["subjects"]
        values = {s["key"]: str((3.1 + i * 0.1) if s["group"] == "Drug" else (4.0 + i * 0.1)) for i, s in enumerate(subs)}
        day = self.read(exp, TODAY, values).get_json()["table"]["stats"]["raw"][0]
        self.assertEqual(day["test"], "Welch's t-test")
        self.assertLess(day["p"], 0.05)
        self.assertEqual(exp_stats.chi_square([(10, 0), (10, 0)])["p"], 1.0)
        xlsx = self.a.get(f"/experiments/{exp}/export.xlsx")
        self.assertEqual(xlsx.status_code, 200)
        from openpyxl import load_workbook
        import io
        book = load_workbook(io.BytesIO(xlsx.data))
        self.assertEqual(book.sheetnames, ["Readout", "Readout, wide", "Manipulations", "Tests by day"])
        self.assertEqual(book["Readout"].max_row, len(subs) + 1)
        csv = self.a.get(f"/experiments/{exp}/export.csv").get_data(as_text=True)
        self.assertIn("Standard length (mm)", csv.splitlines()[0])

    def test_bench_mode_and_its_qr(self):
        exp = self.new("zebrafish")
        html = self.get_ok(self.a, f"/experiments/{exp}/bench")
        self.assertIn("xb-data", html)
        qr = self.a.get(f"/experiments/{exp}/bench.svg")
        self.assertIn(qr.status_code, (200, 404))          # 404 only without segno
        if qr.status_code == 200:
            self.assertIn("<svg", qr.get_data(as_text=True))


class ClutchesAndCohorts(Base):
    def test_larvae_of_a_clutch_and_an_organism_cohort(self):
        from app.db import SessionLocal
        from app.models import ClutchRecord, OrgCohort
        with SessionLocal() as s:
            clutch = ClutchRecord(clutch_id=uniq("CL"), date_of_fertilization=TODAY - timedelta(days=5), larvae_count=60)
            s.add(clutch)
            s.commit()
            clutch_id = clutch.id
        exp = self.new("zebrafish")
        page = self.get_ok(self.a, f"/experiments/{exp}")
        self.assertIn(f'value="clutch:{clutch_id}"', page)
        subs = self.add(exp, "one", f"clutch:{clutch_id}").get_json()["subjects"]
        self.assertEqual((subs[0]["key"], subs[0]["start"]), (f"clutch:{clutch_id}", 60))
        key = self.make_organism_module(self.a)
        mid = self.organism_module_id(key)
        with SessionLocal() as s:
            cohort = OrgCohort(module_id_fk=mid, code=uniq("C"), count_initial=30, count_current=28)
            s.add(cohort)
            s.commit()
            cohort_id = cohort.id
        exp2 = self.new(f"organisms:{key}")
        subs = self.add(exp2, "one", f"cohort:{cohort_id}").get_json()["subjects"]
        self.assertEqual((subs[0]["key"], subs[0]["start"]), (f"cohort:{cohort_id}", 30))
